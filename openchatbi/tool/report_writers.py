"""Convert report content (markdown / CSV / JSON) into Word (.docx) and Excel (.xlsx) files.

Kept dependency-light: the input is the same free-form `content` string the LLM already
passes to `save_report` (usually markdown), so no markdown AST parser is pulled in — a small
line-oriented parser below is enough to get headings, lists, tables, and basic inline
emphasis into a readable Word document, and to lift tabular content into Excel sheets.
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

TABLE_ROW_RE = re.compile(r"^\s*\|(.+)\|\s*$")
TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
NUMBERED_RE = re.compile(r"^\s*\d+[.)]\s+(.*)$")
QUOTE_RE = re.compile(r"^\s*>\s?(.*)$")
FENCE_RE = re.compile(r"^\s*```")
HR_RE = re.compile(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$")


def _split_table_row(line: str) -> list[str]:
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [cell.strip() for cell in inner.split("|")]


def _add_inline_runs(paragraph: Any, text: str) -> None:
    """Split `**bold**`, `*italic*`, and `` `code` `` spans into styled runs."""
    token_re = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")
    for chunk in token_re.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            paragraph.add_run(chunk[1:-1]).italic = True
        else:
            paragraph.add_run(chunk)


def write_docx(file_path: Path, title: str, content: str) -> None:
    """Render markdown-ish `content` into a Word document at `file_path`."""
    from docx import Document

    doc = Document()
    if title:
        doc.add_heading(title, level=0)

    lines = content.splitlines()
    i = 0
    in_code_block = False
    code_lines: list[str] = []
    while i < len(lines):
        line = lines[i]

        if FENCE_RE.match(line):
            if in_code_block:
                code_p = doc.add_paragraph("\n".join(code_lines))
                for run in code_p.runs:
                    run.font.name = "Consolas"
                    run.font.size = None
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        heading_match = HEADING_RE.match(line)
        if heading_match:
            level = min(len(heading_match.group(1)), 9)
            doc.add_heading(heading_match.group(2).strip(), level=level)
            i += 1
            continue

        if HR_RE.match(line):
            doc.add_paragraph("―" * 20)
            i += 1
            continue

        # Markdown table: header row, separator row, then data rows.
        if TABLE_ROW_RE.match(line) and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1]):
            header = _split_table_row(line)
            rows = [header]
            j = i + 2
            while j < len(lines) and TABLE_ROW_RE.match(lines[j]):
                rows.append(_split_table_row(lines[j]))
                j += 1
            table = doc.add_table(rows=len(rows), cols=len(header))
            table.style = "Light Grid Accent 1"
            for r, row_values in enumerate(rows):
                for c, cell_value in enumerate(row_values):
                    if c < len(header):
                        cell = table.cell(r, c)
                        cell.text = cell_value
                        if r == 0:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.bold = True
            i = j
            continue

        quote_match = QUOTE_RE.match(line)
        if quote_match:
            p = doc.add_paragraph()
            p.style = "Intense Quote" if "Intense Quote" in doc.styles else None
            _add_inline_runs(p, quote_match.group(1))
            i += 1
            continue

        bullet_match = BULLET_RE.match(line)
        if bullet_match:
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_runs(p, bullet_match.group(1))
            i += 1
            continue

        numbered_match = NUMBERED_RE.match(line)
        if numbered_match:
            p = doc.add_paragraph(style="List Number")
            _add_inline_runs(p, numbered_match.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        _add_inline_runs(p, line.strip())
        i += 1

    if in_code_block and code_lines:
        code_p = doc.add_paragraph("\n".join(code_lines))
        for run in code_p.runs:
            run.font.name = "Consolas"

    doc.save(str(file_path))


def _extract_markdown_tables(content: str) -> list[tuple[str, list[list[str]]]]:
    """Return `(heading, rows)` pairs for every markdown table found in `content`."""
    lines = content.splitlines()
    tables: list[tuple[str, list[list[str]]]] = []
    last_heading = ""
    i = 0
    while i < len(lines):
        heading_match = HEADING_RE.match(lines[i])
        if heading_match:
            last_heading = heading_match.group(2).strip()
        if TABLE_ROW_RE.match(lines[i]) and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1]):
            rows = [_split_table_row(lines[i])]
            j = i + 2
            while j < len(lines) and TABLE_ROW_RE.match(lines[j]):
                rows.append(_split_table_row(lines[j]))
                j += 1
            tables.append((last_heading, rows))
            i = j
            continue
        i += 1
    return tables


def _looks_like_csv(content: str) -> bool:
    sample = content.strip()
    if not sample or "\n" not in sample:
        return False
    try:
        dialect = csv.Sniffer().sniff(sample.splitlines()[0])
        return dialect.delimiter in {",", ";", "\t"}
    except csv.Error:
        return False


def _style_header_row(ws: Any, ncols: int) -> None:
    from openpyxl.styles import Font, PatternFill

    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")


def _autosize_columns(ws: Any) -> None:
    from openpyxl.utils import get_column_letter

    for col_cells in ws.columns:
        length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, 10), 60)


def write_xlsx(file_path: Path, title: str, content: str) -> None:
    """Render tabular `content` (JSON / CSV / markdown tables / plain text) into an .xlsx file."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    # 1) JSON: list[dict] or dict[str, list] -> single sheet.
    parsed_json: Any = None
    try:
        parsed_json = json.loads(content)
    except (json.JSONDecodeError, TypeError):
        parsed_json = None

    if isinstance(parsed_json, list) and parsed_json and all(isinstance(r, dict) for r in parsed_json):
        ws = wb.create_sheet(title[:31] or "Data")
        headers = list(dict.fromkeys(key for row in parsed_json for key in row.keys()))
        ws.append(headers)
        for row in parsed_json:
            ws.append([row.get(h, "") for h in headers])
        _style_header_row(ws, len(headers))
        _autosize_columns(ws)
        wb.save(str(file_path))
        return

    if isinstance(parsed_json, dict) and parsed_json and all(isinstance(v, list) for v in parsed_json.values()):
        ws = wb.create_sheet(title[:31] or "Data")
        headers = list(parsed_json.keys())
        ws.append(headers)
        max_len = max(len(v) for v in parsed_json.values())
        for idx in range(max_len):
            ws.append([parsed_json[h][idx] if idx < len(parsed_json[h]) else "" for h in headers])
        _style_header_row(ws, len(headers))
        _autosize_columns(ws)
        wb.save(str(file_path))
        return

    # 2) Markdown tables -> one sheet per table.
    md_tables = _extract_markdown_tables(content)
    if md_tables:
        for idx, (heading, rows) in enumerate(md_tables, start=1):
            sheet_name = (heading or f"Table{idx}")[:31] or f"Table{idx}"
            base_name, suffix = sheet_name, 1
            while sheet_name in wb.sheetnames:
                suffix += 1
                sheet_name = f"{base_name[:28]}_{suffix}"
            ws = wb.create_sheet(sheet_name)
            for row in rows:
                ws.append(row)
            if rows:
                _style_header_row(ws, len(rows[0]))
            _autosize_columns(ws)
        wb.save(str(file_path))
        return

    # 3) Raw CSV content -> single sheet.
    if _looks_like_csv(content):
        ws = wb.create_sheet(title[:31] or "Data")
        for row in csv.reader(io.StringIO(content)):
            ws.append(row)
        if ws.max_row >= 1:
            _style_header_row(ws, ws.max_column)
        _autosize_columns(ws)
        wb.save(str(file_path))
        return

    # 4) Fallback: one line of plain text per row.
    ws = wb.create_sheet(title[:31] or "Report")
    ws.append([title or "Report"])
    _style_header_row(ws, 1)
    for line in content.splitlines() or [content]:
        ws.append([line])
    _autosize_columns(ws)
    wb.save(str(file_path))

"""Tests for the `save_report` tool, including the docx/xlsx Office export path."""

from unittest.mock import patch

import pytest
from docx import Document
from openpyxl import load_workbook

from openchatbi.tool.save_report import save_report

MARKDOWN_CONTENT = """# Sales Report

## Overview
Revenue grew **12%** month over month.

- East: up 20%
- South: up 5%

## Details

| Region | Revenue | YoY |
| --- | --- | --- |
| East | 120000 | 20% |
| South | 80000 | 5% |
"""


@pytest.fixture
def report_dir(tmp_path):
    with patch("openchatbi.tool.save_report.config.get") as mock_get:
        mock_get.return_value.report_directory = str(tmp_path)
        yield tmp_path


class TestSaveReportTextFormats:
    def test_save_markdown_report(self, report_dir):
        result = save_report.invoke({"content": "# Hello\nWorld", "title": "My Report", "file_format": "md"})
        assert "Report saved successfully!" in result
        assert "Download link: /api/download/report/" in result
        saved_files = list(report_dir.glob("*My_Report.md"))
        assert len(saved_files) == 1
        assert saved_files[0].read_text(encoding="utf-8") == "# Hello\nWorld"

    def test_unsupported_format_raises(self, report_dir):
        with pytest.raises(ValueError, match="Unsupported file format"):
            save_report.invoke({"content": "x", "title": "t", "file_format": "pdf"})


class TestSaveReportDocx:
    def test_save_docx_report(self, report_dir):
        result = save_report.invoke({"content": MARKDOWN_CONTENT, "title": "Sales Report", "file_format": "docx"})
        assert "Report saved successfully!" in result

        saved_files = list(report_dir.glob("*.docx"))
        assert len(saved_files) == 1

        doc = Document(str(saved_files[0]))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert "Overview" in headings
        assert "Details" in headings

        bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
        assert any("East" in b for b in bullets)

        assert len(doc.tables) == 1
        table_rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
        assert table_rows[0] == ["Region", "Revenue", "YoY"]
        assert ["East", "120000", "20%"] in table_rows


class TestSaveReportXlsx:
    def test_save_xlsx_from_markdown_table(self, report_dir):
        result = save_report.invoke({"content": MARKDOWN_CONTENT, "title": "Sales Report", "file_format": "xlsx"})
        assert "Report saved successfully!" in result

        saved_files = list(report_dir.glob("*.xlsx"))
        assert len(saved_files) == 1

        wb = load_workbook(str(saved_files[0]))
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0] == ("Region", "Revenue", "YoY")
        assert ("East", "120000", "20%") in rows

    def test_save_xlsx_from_json_list(self, report_dir):
        import json

        content = json.dumps([{"name": "Alice", "score": 90}, {"name": "Bob", "score": 85}])
        result = save_report.invoke({"content": content, "title": "Scores", "file_format": "xlsx"})
        assert "Report saved successfully!" in result

        saved_files = list(report_dir.glob("*.xlsx"))
        assert len(saved_files) == 1
        wb = load_workbook(str(saved_files[0]))
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0] == ("name", "score")
        assert ("Alice", 90) in rows

    def test_save_xlsx_from_csv(self, report_dir):
        content = "name,score\nAlice,90\nBob,85\n"
        result = save_report.invoke({"content": content, "title": "Scores", "file_format": "xlsx"})
        assert "Report saved successfully!" in result

        saved_files = list(report_dir.glob("*.xlsx"))
        wb = load_workbook(str(saved_files[0]))
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0] == ("name", "score")
        assert ("Alice", "90") in rows
